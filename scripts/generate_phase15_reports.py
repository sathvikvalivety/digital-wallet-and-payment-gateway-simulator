import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import datetime
import os

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def create_hardening_checklist():
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin, s.bottom_margin = Inches(1.0), Inches(1.0)
        s.left_margin, s.right_margin = Inches(1.0), Inches(1.0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("PHASE 15: INFRASTRUCTURE & APPLICATION HARDENING CHECKLIST\nDWPG SIMULATOR SYSTEM HARDENING AUDIT")
    r.font.name, r.font.size, r.font.bold = "Arial", Pt(18), True
    r.font.color.rgb = RGBColor(30, 58, 138)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = sub.add_run("Course: 24CYS401 Secure Software Engineering | Examination Session 2026\nLead DevSecOps Architect: Sathvik Valivety (GitHub: sathvikvalivety)")
    r_sub.font.size, r_sub.font.italic = Pt(10), True

    doc.add_heading("1. Multi-Tier Hardening Verification Matrix", level=1)
    
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0]
    for i, h in enumerate(["Layer", "Hardening Control Specification", "Implementation Evidence", "Audit Verification"]):
        hdr.cells[i].paragraphs[0].add_run(h).font.bold = True
        set_cell_background(hdr.cells[i], "1E3A8A")
        hdr.cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

    items = [
        ("Operating System", "Non-root execution (UID 10001, GID 10001) for all processes", "Alpine user appuser created in Dockerfile", "PASSED"),
        ("Operating System", "Minimal Linux distribution stripping compilers & package tools", "eclipse-temurin:21-jre-alpine & nginx-unprivileged", "PASSED"),
        ("JVM Runtime", "Container-aware heap allocation (-XX:MaxRAMPercentage=75.0)", "Configured in backend Dockerfile ENTRYPOINT", "PASSED"),
        ("JVM Runtime", "CSPRNG non-blocking entropy source (-Djava.security.egd)", "Mapped to /dev/urandom in container startup", "PASSED"),
        ("Spring Framework", "Suppression of internal stacktraces & exception messages", "application.yml include-stacktrace: never", "PASSED"),
        ("Spring Framework", "Strict Actuator endpoint exposure & detail authorization", "management.endpoints.web.exposure.include: health,info,metrics", "PASSED"),
        ("Spring Security", "Stateless JWT authentication with BCrypt password hashing", "SecurityConfig filter chain with BCryptPasswordEncoder (strength 12)", "PASSED"),
        ("Spring Security", "Role-Based Access Control on sensitive endpoints (@PreAuthorize)", "Enforced across Wallet, Merchant, Payment & Admin controllers", "PASSED"),
        ("Database", "Principle of least privilege non-root database credentials", "Configured via MARIADB_USER dwpg_user in k8s/02-secret.yaml", "PASSED"),
        ("Database", "ACID pessimistic row locking preventing concurrency hazards", "WalletRepository @Lock(PESSIMISTIC_WRITE) on balance mutations", "PASSED"),
        ("Web Proxy", "Defensive HTTP Security Headers (CSP, X-Frame-Options, XSS)", "Injected via frontend/nginx.conf reverse proxy", "PASSED"),
        ("Kubernetes", "Pod Security Standards (Restricted Profile) compliance", "runAsNonRoot: true, allowPrivilegeEscalation: false, drop: [ALL]", "PASSED"),
        ("Kubernetes", "Zero-Trust NetworkPolicy isolating database to backend pods", "k8s/10-networkpolicy.yaml default-deny ingress on port 3306", "PASSED")
    ]

    for cat, ctrl, ev, audit in items:
        row = table.add_row()
        row.cells[0].paragraphs[0].add_run(cat).font.bold = True
        row.cells[0].paragraphs[0].runs[0].font.size = Pt(8.5)
        row.cells[1].paragraphs[0].add_run(ctrl).font.size = Pt(8.5)
        row.cells[2].paragraphs[0].add_run(ev).font.size = Pt(8.5)
        r_a = row.cells[3].paragraphs[0].add_run(audit)
        r_a.font.bold = True
        r_a.font.size = Pt(8.5)
        r_a.font.color.rgb = RGBColor(22, 101, 52)
        for c in row.cells:
            set_cell_background(c, "F8FAFC")

    doc.save("docs/15-hardening/hardening-checklist.docx")
    print("Hardening checklist created.")

def create_secure_deployment_checklist():
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin, s.bottom_margin = Inches(1.0), Inches(1.0)
        s.left_margin, s.right_margin = Inches(1.0), Inches(1.0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("PHASE 15: SECURE DEPLOYMENT GOVERNANCE CHECKLIST\nPRODUCTION READINESS & RELEASE CRITERIA")
    r.font.name, r.font.size, r.font.bold = "Arial", Pt(18), True
    r.font.color.rgb = RGBColor(30, 58, 138)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = sub.add_run("Course: 24CYS401 Secure Software Engineering | Examination Session 2026\nLead DevSecOps Architect: Sathvik Valivety (GitHub: sathvikvalivety)")
    r_sub.font.size, r_sub.font.italic = Pt(10), True

    doc.add_heading("1. Production Release Gate Verification", level=1)

    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0]
    for i, h in enumerate(["Checkpoint Phase", "Mandatory Security Standard", "Verification Artifact", "Release Sign-Off"]):
        hdr.cells[i].paragraphs[0].add_run(h).font.bold = True
        set_cell_background(hdr.cells[i], "1E3A8A")
        hdr.cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

    checks = [
        ("Code Quality Gate", "SonarQube SAST status = OK, 0 Blocker/Critical vulnerabilities", "docs/11-secure-build/sonarqube-report.docx", "APPROVED"),
        ("Regression Testing", "25 automated JUnit tests pass (concurrency, fuzzing, BOLA)", "backend/src/test/java test suite logs", "APPROVED"),
        ("Secret Hygiene", "Zero plain text secrets or passwords committed to Git", ".gitignore enforcement, .env.example template", "APPROVED"),
        ("Container Integrity", "Multi-stage non-root container images built & verified", "Docker images (UID 10001 & UID 101 verified)", "APPROVED"),
        ("Vulnerability Scan", "Trivy container scan shows zero unpatched critical CVEs", "CI/CD Pipeline Stage 5 Trivy task", "APPROVED"),
        ("Orchestration IaC", "Kubernetes manifests pass schema validation & PSS Restricted", "k8s/ manifests validated via dry-run/kubeconform", "APPROVED"),
        ("Zero-Trust Network", "Database traffic restricted exclusively to backend pods", "k8s/10-networkpolicy.yaml NetworkPolicy", "APPROVED"),
        ("Health & Liveness", "Automated container probes configured for rolling update", "livenessProbe & readinessProbe on port 8080/3000", "APPROVED"),
        ("Audit Compliance", "Immutable audit logging capturing all security-relevant events", "AuditService recording auth, payments, replays", "APPROVED")
    ]

    for ph, std, art, sign in checks:
        row = table.add_row()
        row.cells[0].paragraphs[0].add_run(ph).font.bold = True
        row.cells[0].paragraphs[0].runs[0].font.size = Pt(8.5)
        row.cells[1].paragraphs[0].add_run(std).font.size = Pt(8.5)
        row.cells[2].paragraphs[0].add_run(art).font.size = Pt(8.5)
        r_s = row.cells[3].paragraphs[0].add_run(sign)
        r_s.font.bold = True
        r_s.font.size = Pt(8.5)
        r_s.font.color.rgb = RGBColor(22, 101, 52)
        for c in row.cells:
            set_cell_background(c, "F8FAFC")

    doc.save("docs/15-hardening/secure-deployment-checklist.docx")
    print("Secure deployment checklist created.")

def create_logging_monitoring_plan():
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin, s.bottom_margin = Inches(1.0), Inches(1.0)
        s.left_margin, s.right_margin = Inches(1.0), Inches(1.0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("PHASE 15: LOGGING, MONITORING AND SIEM INTEGRATION PLAN\nSECURITY INCIDENT & OBSERVABILITY ARCHITECTURE")
    r.font.name, r.font.size, r.font.bold = "Arial", Pt(18), True
    r.font.color.rgb = RGBColor(30, 58, 138)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = sub.add_run("Course: 24CYS401 Secure Software Engineering | Examination Session 2026\nLead DevSecOps Architect: Sathvik Valivety (GitHub: sathvikvalivety)")
    r_sub.font.size, r_sub.font.italic = Pt(10), True

    doc.add_heading("1. Observability & Security Logging Philosophy", level=1)
    doc.add_paragraph(
        "Financial systems mandate non-repudiable audit trails and continuous observability. "
        "The DWPG simulator incorporates structured security audit logging, PII scrubbing (PCI-DSS compliance), "
        "correlation ID tracing, and SIEM alert detection rules to identify attacks in real-time."
    )

    doc.add_heading("2. Immutable Security Audit Trail (AuditService)", level=1)
    doc.add_paragraph(
        "Security-relevant operations generate immutable entries in the 'audit_logs' table containing:\n"
        "• Timestamp: ISO-8601 UTC timestamp\n"
        "• Event Type: AUTH_LOGIN_SUCCESS, AUTH_LOGIN_FAILED, USER_REGISTERED, PAYMENT_INITIATED, "
        "PAYMENT_CONFIRMED, PAYMENT_REFUNDED, REPLAY_DETECTED, AUTHORIZATION_FAILURE, SUSPICIOUS_ACTIVITY\n"
        "• Actor Username & User ID: Authenticated principal initiating the action\n"
        "• Resource Type & ID: Target wallet, payment, or merchant identifier\n"
        "• Outcome: SUCCESS or FAILED\n"
        "• Client IP Address: Extracted from X-Forwarded-For or RemoteAddr\n"
        "• Correlation ID: Distributed request trace identifier"
    )

    doc.add_heading("3. Sensitive Data Scrubbing & Masking Engine", level=1)
    doc.add_paragraph(
        "To prevent secret leakage in logs, AuditService executes pre-storage regex scrubbing:\n"
        "1. Password / Secret Scrubbing: Matches any JSON keys 'password', 'secret', 'token' and replaces values with '***REDACTED***'.\n"
        "2. PAN / Payment Card Masking: Matches Visa and MasterCard Luhn patterns and replaces digits with '****-****-****-****'."
    )

    doc.add_heading("4. SIEM Threat Detection Rules & Alert Thresholds", level=1)
    
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0]
    for i, h in enumerate(["Threat Scenario", "SIEM Detection Rule", "Severity", "Automated Response"]):
        hdr.cells[i].paragraphs[0].add_run(h).font.bold = True
        set_cell_background(hdr.cells[i], "1E3A8A")
        hdr.cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

    alerts = [
        ("Brute-Force Authentication", ">= 5 AUTH_LOGIN_FAILED within 60s for same IP or username", "HIGH", "Trigger account temporary lockout & IP throttle alert"),
        ("Replay Attack Attempt", "Event REPLAY_DETECTED logged by IdempotencyService", "CRITICAL", "Reject request with HTTP 400 & flag client IP for SOC investigation"),
        ("Double Spending Race", "Repeated InsufficientFundsException in concurrent window", "HIGH", "Record fraud warning & alert compliance team"),
        ("BOLA / IDOR Violation", "Event AUTHORIZATION_FAILURE logged on foreign wallet access", "CRITICAL", "Revoke active JWT token & log tenant breach attempt")
    ]

    for scn, rule, sev, resp in alerts:
        row = table.add_row()
        row.cells[0].paragraphs[0].add_run(scn).font.bold = True
        row.cells[0].paragraphs[0].runs[0].font.size = Pt(8.5)
        row.cells[1].paragraphs[0].add_run(rule).font.size = Pt(8.5)
        r_sev = row.cells[2].paragraphs[0].add_run(sev)
        r_sev.font.bold = True
        r_sev.font.size = Pt(8.5)
        if sev == "CRITICAL":
            r_sev.font.color.rgb = RGBColor(220, 38, 38)
        else:
            r_sev.font.color.rgb = RGBColor(217, 119, 6)
        row.cells[3].paragraphs[0].add_run(resp).font.size = Pt(8.5)
        for c in row.cells:
            set_cell_background(c, "F8FAFC")

    doc.add_paragraph()

    # Section 5
    doc.add_heading("5. Metric Telemetry & Dashboards", level=1)
    doc.add_paragraph(
        "System telemetry is scraped by Prometheus via Spring Boot Actuator endpoint /actuator/metrics:\n"
        "• jvm.memory.used & jvm.memory.max (Heap utilization)\n"
        "• hikaricp.connections.active & hikaricp.connections.idle (Database connection pool pressure)\n"
        "• http.server.requests (Latency percentiles p50, p95, p99 and HTTP status codes)\n"
        "• dwpg.payments.processed.count (Business transaction velocity)."
    )

    doc.save("docs/15-hardening/logging-monitoring-plan.docx")
    print("Logging & monitoring plan created.")

if __name__ == "__main__":
    os.makedirs("docs/15-hardening", exist_ok=True)
    create_hardening_checklist()
    create_secure_deployment_checklist()
    create_logging_monitoring_plan()
