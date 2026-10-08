import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import os

def create_traceability_matrix():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "RTM Traceability Matrix"

    # Styling
    font_hdr = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    fill_hdr = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    font_row = Font(name="Calibri", size=9.5)
    font_bold = Font(name="Calibri", size=9.5, bold=True)
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )
    fill_even = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_status = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    font_status = Font(name="Calibri", size=9.5, bold=True, color="166534")

    headers = [
        "Req ID", "Category", "Requirement Description", "SRS Ref", 
        "UML Artifact", "DFD / Trust Boundary", "Architecture Component",
        "STRIDE / Attack Tree", "Jira Story", "Code Implementation",
        "Test Verification", "Hardening / Quality Gate", "Compliance"
    ]

    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = font_hdr
        cell.fill = fill_hdr
        cell.alignment = align_center

    rows = [
        ("FR-001", "Functional", "User Registration with Argon2/BCrypt hashing", "SRS §3.1", "UC-REG-001", "DFD-1 Process 1.0 (Public Zone)", "AuthService / User", "Spoofing (T-01)", "DWPG-10", "AuthService.register()", "AuthServiceTest.testRegisterUserSuccess", "SonarQube S6437 check", "VERIFIED"),
        ("FR-002", "Functional", "User Authentication with secure JWT issuance", "SRS §3.2", "UC-AUTH-001", "DFD-1 Process 1.0 (DMZ)", "JwtTokenProvider / SecurityConfig", "Spoofing / Repudiation", "DWPG-11", "AuthService.login()", "AuthServiceTest.testLoginSuccess", "BCrypt 12 rounds", "VERIFIED"),
        ("FR-003", "Functional", "Simulated Wallet Creation upon registration", "SRS §3.3", "UC-WAL-001", "DFD-1 Process 2.0 (Trusted Internal)", "WalletService / Wallet", "Elevation of Privilege", "DWPG-12", "WalletService.createWallet()", "WalletServiceTest.testCreateWalletSuccess", "Zero initial balance invariant", "VERIFIED"),
        ("FR-004", "Functional", "Adding simulated funds (Top-up)", "SRS §3.4", "UC-WAL-002", "DFD-1 Process 2.0 (Vault Zone)", "WalletService / TransactionService", "Tampering / Denial of Service", "DWPG-13", "WalletService.topUp()", "WalletServiceTest.testTopUpSuccess", "Atomic ACID transaction", "VERIFIED"),
        ("FR-005", "Functional", "Merchant Registration with 256-bit API key", "SRS §3.5", "UC-MER-001", "DFD-1 Process 3.0 (Merchant Zone)", "MerchantService / Merchant", "Elevation of Privilege", "DWPG-14", "MerchantService.registerMerchant()", "PaymentSecurityIntegrationTest.testMerchantContext", "Role ROLE_MERCHANT guard", "VERIFIED"),
        ("FR-006", "Functional", "Payment Initiation with mandatory Idempotency-Key", "SRS §3.6", "UC-PAY-001", "DFD-1 Process 4.0 (Payment Gateway)", "PaymentService / IdempotencyService", "Replay (T-03) / Attack Node 1.2", "DWPG-15 / DWPG-19", "PaymentService.processPayment()", "PaymentSecurityIntegrationTest.testIdempotency", "SHA-256 payload digest cache", "VERIFIED"),
        ("FR-007", "Functional", "Payment Confirmation & state machine transitions", "SRS §3.7", "UC-PAY-002", "DFD-1 Process 4.0 (Core Engine)", "Payment / State Machine", "Tampering / State Inversion", "DWPG-16", "Payment.setStatus(CONFIRMED)", "PaymentSecurityIntegrationTest.testPaymentLifecycle", "One-way status transition rule", "VERIFIED"),
        ("FR-008", "Functional", "Refund Processing & original transaction linkage", "SRS §3.8", "UC-REF-001", "DFD-1 Process 5.0 (Refund Rail)", "RefundService / Refund", "Double Spending / Repeat Refund", "DWPG-17", "RefundService.processRefund()", "PaymentSecurityIntegrationTest.testDuplicateRefundRejected", "Payment REFUNDED status check", "VERIFIED"),
        ("FR-009", "Functional", "Transaction History & Tamper-Evident Ledger", "SRS §3.9", "UC-LED-001", "DFD-1 Process 6.0 (Audit Vault)", "TransactionService / Ledger", "Tampering (T-02) / Attack Node 3.1", "DWPG-18", "TransactionService.getMyTransactions()", "PaymentSecurityIntegrationTest.testTamperHash", "SHA-256 state hash validation", "VERIFIED"),
        ("FR-010", "Functional", "Administrative Audit Log Explorer & SIEM Feed", "SRS §3.10", "UC-ADM-001", "DFD-1 Process 7.0 (Admin Console)", "AuditService / AuditLog", "Repudiation / Info Disclosure", "DWPG-21", "AuditService.getAllAuditLogs()", "PaymentSecurityIntegrationTest.testAdminAudit", "ROLE_ADMIN authorization check", "VERIFIED"),
        ("FR-011", "Functional", "Anti-Replay Attack Idempotency Filter", "SRS §3.11", "UC-SEC-001", "DFD-1 Process 4.1 (Replay Filter)", "IdempotencyService / Record", "Replay Attack (SEC-004)", "DWPG-20", "IdempotencyService.checkIdempotency()", "PaymentSecurityIntegrationTest.testReplayAttackRejected", "Duplicate key rejection + caching", "VERIFIED"),
        ("FR-012", "Functional", "BOLA Tenant Isolation & Ownership Enforcement", "SRS §3.12", "UC-SEC-002", "DFD-1 Process 2.1 (BOLA Guard)", "WalletService / SecurityContext", "Elevation of Privilege (SEC-007)", "DWPG-22", "WalletService.getWalletById()", "WalletServiceTest.testGetWalletByIdForeignUserAccessDenied", "SecurityContext username match", "VERIFIED"),
        ("NFR-001", "Non-Functional", "Sub-200ms Transaction Processing Latency", "SRS §4.1", "Analysis Model", "Trust Boundary 2", "HikariCP / Striped Lock", "Denial of Service", "DWPG-19", "PaymentService concurrent execution", "PaymentConcurrencyTest (10 Threads < 2s)", "Hikari pool size 30", "VERIFIED"),
        ("NFR-002", "Non-Functional", "ACID Financial Isolation & Consistency", "SRS §4.2", "Analysis Model", "Database Tier", "TransactionTemplate / MariaDB", "Tampering / Inconsistency", "DWPG-19", "WalletRepository.findByUserIdForUpdate()", "PaymentConcurrencyTest (Zero Overdraft)", "SELECT FOR UPDATE row lock", "VERIFIED"),
        ("NFR-003", "Non-Functional", "99.95% Availability with Graceful Degradation", "SRS §4.3", "Component Diagram", "Container Tier", "Kubernetes 2-Replica Deployment", "Denial of Service", "DWPG-8", "backend-deployment.yaml replicas: 2", "k8s readiness/liveness probes", "Rolling update strategy", "VERIFIED"),
        ("NFR-004", "Non-Functional", "PCI-DSS Sensitive Data Masking & Redaction", "SRS §4.4", "Component Diagram", "Audit Tier", "AuditService regex scrubber", "Information Disclosure", "DWPG-18", "AuditService.scrubSensitiveData()", "AuditService regex validation", "Redaction of PAN & passwords", "VERIFIED"),
        ("NFR-005", "Non-Functional", "REST API Conformance & OpenAPI Schema", "SRS §4.5", "Component Diagram", "REST Controller Layer", "GlobalExceptionHandler / DTOs", "Improper Input Handling", "DWPG-10..18", "DWPG REST Controllers", "PaymentInputFuzzingTest (12 Vectors)", "RFC 7807 Error Responses", "VERIFIED"),
        ("NFR-006", "Non-Functional", "Scalable Zero-Downtime Container Deployment", "SRS §4.6", "Component Diagram", "Kubernetes Pods", "Docker / Minikube", "Denial of Service", "DWPG-8", "k8s/ manifests & docker-compose", "Docker non-root execution (UID 10001)", "PSS Restricted compliance", "VERIFIED"),
        ("SEC-001", "Security", "BCrypt Hashing with Zero Plaintext Credentials", "SRS §5.1", "UC-AUTH-001", "Public -> DMZ", "SecurityConfig / PasswordEncoder", "Information Disclosure / Spoofing", "DWPG-10", "BCryptPasswordEncoder(12)", "AuthServiceTest.testRegisterHash", "SonarQube S6437 check (0 findings)", "VERIFIED"),
        ("SEC-002", "Security", "Stateless JWT with HMAC-SHA256 256-bit Key", "SRS §5.2", "UC-AUTH-002", "DMZ -> Internal", "JwtTokenProvider", "Tampering / Spoofing", "DWPG-11", "JwtTokenProvider.generateToken()", "AuthServiceTest.testLoginTokenValid", "Signed claims + expiration check", "VERIFIED"),
        ("SEC-003", "Security", "Role-Based Access Control (USER, MERCHANT, ADMIN)", "SRS §5.3", "Analysis Model", "Internal Boundary", "Spring Security @PreAuthorize", "Elevation of Privilege", "DWPG-11..22", "Controller method security annotations", "PaymentSecurityIntegrationTest.testRbacEnforcement", "403 Forbidden on role mismatch", "VERIFIED"),
        ("SEC-004", "Security", "Mandatory Idempotency-Key on Mutating Endpoints", "SRS §5.4", "UC-SEC-001", "Gateway Boundary", "IdempotencyService", "Replay Attacks (CWE-294)", "DWPG-15 / DWPG-20", "IdempotencyService.checkIdempotency()", "PaymentSecurityIntegrationTest.testIdempotency", "Cached responses on identical retry", "VERIFIED"),
        ("SEC-005", "Security", "Replay Attack Detection & Payload Tampering Guard", "SRS §5.5", "UC-SEC-001", "Gateway Boundary", "IdempotencyService SHA-256", "Replay / Tampering (CWE-799)", "DWPG-20", "IdempotencyRecord requestHash compare", "PaymentSecurityIntegrationTest.testReplayAttackRejected", "REPLAY_DETECTED event logged", "VERIFIED"),
        ("SEC-006", "Security", "Race Condition & Double-Spending Prevention", "SRS §5.6", "UC-PAY-001", "Vault Tier", "Striped User Lock + Pessimistic Row Lock", "Race Conditions (CWE-362)", "DEF-001 / DWPG-19", "WalletRepository.findByUserIdForUpdate()", "PaymentConcurrencyTest (10 Parallel Threads)", "Zero double spending confirmed", "VERIFIED"),
        ("SEC-007", "Security", "BOLA / IDOR Defense on Wallets & Transactions", "SRS §5.7", "UC-SEC-002", "Internal Tier", "WalletService ownership check", "Broken Object Level Auth (CWE-639)", "DWPG-22", "WalletService.getWalletById()", "WalletServiceTest.testGetWalletByIdForeignUserAccessDenied", "403 UnauthorizedAccessException", "VERIFIED"),
        ("SEC-008", "Security", "Boundary Fuzzing & SQL/XSS Injection Immunity", "SRS §5.8", "Analysis Model", "Perimeter Boundary", "Jakarta Validation + Parameterized JPA", "Injection (CWE-89, CWE-79)", "DWPG-15", "PaymentRequest / GlobalExceptionHandler", "PaymentInputFuzzingTest (12 Attack Vectors)", "100% boundary tests pass", "VERIFIED"),
        ("SEC-009", "Security", "Immutable Security Audit Logging & PII Scrubbing", "SRS §5.9", "UC-ADM-001", "Audit Tier", "AuditService / AuditLog table", "Repudiation (CWE-778)", "DWPG-18", "AuditService.logEvent()", "PaymentSecurityIntegrationTest.testAdminAudit", "PAN and password auto-redacted", "VERIFIED"),
        ("SEC-010", "Security", "Zero-Trust Network Segmentation & Hardened Containers", "SRS §5.10", "Component Diagram", "Network Layer", "Kubernetes NetworkPolicy & Dockerfile", "Container Escape / Lateral Movement", "DWPG-8", "k8s/10-networkpolicy.yaml", "Docker non-root UID 10001 verification", "Port 3306 locked to backend pods", "VERIFIED")
    ]

    for row_idx, rdata in enumerate(rows, start=2):
        ws.append(rdata)
        fill = fill_even if row_idx % 2 == 0 else PatternFill(fill_type=None)
        for col_idx in range(1, len(rdata) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = font_row
            cell.border = thin_border
            if fill.fill_type:
                cell.fill = fill
            if col_idx in [1, 2, 4, 9, 13]:
                cell.alignment = align_center
            else:
                cell.alignment = align_left
            if col_idx == 13:
                cell.fill = fill_status
                cell.font = font_status

    # Column Widths
    col_widths = {1: 10, 2: 14, 3: 34, 4: 12, 5: 14, 6: 25, 7: 25, 8: 24, 9: 15, 10: 28, 11: 30, 12: 24, 13: 14}
    for c_idx, width in col_widths.items():
        ws.column_dimensions[openpyxl.utils.get_column_letter(c_idx)].width = width

    os.makedirs("docs/final", exist_ok=True)
    wb.save("docs/final/traceability-matrix.xlsx")
    print("Traceability matrix saved to docs/final/traceability-matrix.xlsx")

if __name__ == "__main__":
    create_traceability_matrix()
